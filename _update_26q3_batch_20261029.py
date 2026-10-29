# -*- coding: utf-8 -*-
import copy
import datetime
import json
import shutil
from pathlib import Path

import openpyxl
from openpyxl.formula.translate import Translator

ROOT = Path(r'D:/Kimi/金属总网/网站构建/财报汇总')
TODAY = '2026-10-29'
FILES = {
    '铜': ROOT / '铜/全球铜企季度产量梳理.xlsx',
    '镍': ROOT / '镍/全球镍企季度产量梳理.xlsx',
    '锂': ROOT / '锂/全球锂企季度产量梳理.xlsx',
    '锡': ROOT / '锡/海外主要公司产量.xlsx',
}

DATA = {
'铜': [
    ('铜矿·季度产量', 'FreePort', None, 832 * 0.4536, '26Q3铜产量832M lbs（并表可回收）；Grasberg矿区磨矿处理量恢复至正常产能67%，PTFI冶炼厂8月底重启'),
    ('铜矿·季度产量', 'Antofagasta', None, 151.9, '26Q3铜151.9kt（环比+7%）：Los Pelambres恢复并消化上季库存；全年指引维持'),
    ('铜矿·季度产量', 'RioTinto', None, 205.0, '26Q3 consolidated铜205kt（同比持平/环比-4%）：OT地下矿同比+27%，Kennecott闪速转炉8月底完成重建'),
    ('铜矿·季度产量', 'Nornickel', '本部', 89.704, '26Q3集团铜108.704kt，其中外贝加尔19kt，本部89.704kt；集团9M铜311.777kt（-0.5%）'),
    ('铜矿·季度产量', 'Nornickel', 'Bystrinsky', 19.0, '26Q3外贝加尔铜精矿含铜19kt（环比+5%），矿石量、品位与回收率提升'),
    ('铜矿·季度产量', 'South32', None, 18.7, 'FY27Q1 Sierra Gorda 45%应付铜18.7kt（环比+17%），铜品位0.41%→0.46%'),
    ('铜矿·季度产量', 'BHP', None, 400.7, 'FY27Q1集团铜400.7kt（同比-19%）：计划品位下降与检修，叠加智利暴雪、安全停产及Carrapateena输送带故障'),
    ('铜矿·季度产量', 'First Quantum', None, 108.959, '26Q3铜108.959kt（Q2=100.487）：Kansanshi S3强劲，Cobre Panamá三条产线季内陆续恢复运行'),
    ('铜矿·季度产量', 'SouthernCopper', None, 232.396, '26Q3矿产铜232.396kt（环比-1.1%）：秘鲁品位下降，墨西哥产量+2.5%部分对冲'),
    ('铜矿·季度产量', 'Anglo American', None, 174.9, '26Q3铜174.9kt（同比-5%）：智利创纪录降雪及非计划检修，Collahuasi处理低品位库存矿'),
    ('铜矿·季度产量', 'Sandfire Resources', None, 38.9, 'FY27Q1铜当量38.9kt：MATSA 22.7+Motheo 16.1kt，达全年指引中值25%'),
    ('铜矿·季度产量', 'Ivanhoe Mines', None, round(76.401 * 0.396, 4), '26Q3 Kamoa-Kakula 100%铜76.401kt，按39.6%权益=30.2548kt；含阳极铜61.506kt'),
    ('国内铜企·季度产量', 'MMG', None, 13.1215, '26Q3集团铜131,215吨（同比+3%），折13.1215万吨（100%基准）'),
],
'镍': [
    ('一级镍·季度产量', 'Nornickel', None, 61103, '26Q3镍61,103吨（环比+41%）：处理Norilsk航运间歇期积累原料，9M=146,131吨（+4%）'),
    ('镍铁NPI·季度产量', 'Anglo American', None, 10400, '26Q3 Barro Alto+Codemin镍铁10.4kt（同比+3%），检修时间差异；镍业务出售MMG仍待欧盟审批'),
    ('镍铁NPI·季度产量', 'Antam', None, 3850, '26Q3镍铁3,850 TNi，9M=11,638 TNi；9M销量13,971 TNi（+71%）'),
    ('镍铁NPI·季度产量', 'Nickel Industries', None, 27822, '26Q3 RKEF镍27,822吨（NPI 239,959吨、品位11.6%）；干旱制约运行，10月约75%铭牌负荷'),
],
'锂': [
    ('锂资源·季度产量', 'Pilbara Minerals', None, 267600 * 0.125, 'FY27Q1 Pilgangoora精矿267,600dmt（环比+25%创纪录）×0.125=33,450吨LCE；P2000扩建获批'),
    ('锂资源·季度产量', 'Mineral Resources', None, (102000 + 71000) * 0.125, 'FY27Q1 Wodgina 102k+Mt Marion 71k dmt，100%口径×0.125=21,625吨LCE；Bald Hill 27k不计入本行既有口径'),
    ('锂资源·季度产量', 'IGO', None, 435000 * 0.125, 'FY27Q1 Greenbushes精矿435kt（环比+12%创纪录）×0.125=54,375吨LCE，CGP3达铭牌产能'),
    ('锂资源·季度产量', 'Rio Tinto（历史Arcadium）', None, 13000, '26Q3 LCE 13.0kt（同比+4%/环比-11%）：恶劣天气影响Q3，Fenix 1B继续爬坡；9M=40.3kt（+33%）'),
    ('锂资源·季度产量', 'Core Lithium', None, '已复产首批，吨位待季报', 'Finniss 9/8产出重启后首批锂精矿，Q4目标首批发运；截至本次核查未见官方Q3吨位'),
],
'锡': [
    ('锡矿·季度产量', 'Alphamin', None, 5030, '26Q3含锡5,030吨（Q2=5,013）：处理量213,454吨、品位3.2%、回收率72.9%；Q4雨季物流风险已增备卡车与燃料'),
    ('锡矿·季度产量', 'Metals X', None, 2663, '26Q3 Renison 100%锡精矿含锡2,663吨（Q2=2,809）；Metals X 50%权益=1,332吨'),
    ('锡矿·季度产量', 'Andrada Mining', None, 290, 'FY27Q2（截至8月）Uis含锡290吨（同比+6%）；H1=576吨（+13%），破碎与矿石分选改造继续'),
],
}

PENDING = {
    'Teck': '待发布(2026-10-29)', 'Glencore': '待发布(2026-10-29)', 'Boliden': '待发布(2026-10-29)',
    'Nexa': '待发布(2026-11上旬)', 'NEXA': '待发布(2026-11上旬)', 'KGHM': '待发布(预计2026-11中旬)',
    'Vale': '待发布(预计2026-10月末)', 'Vale(PTVI)': '待发布(预计2026-10月末)',
    'Eramet': '待发布(2026-10-29)', 'Merdeka(MBMA)': '待发布(预计2026-11上旬)',
    'Harita Nickel (TBP)': '待发布(预计2026-11上旬)', 'Capstone Copper': '待发布(2026-10-29美股盘后)',
    'Lundin Mining': '待发布(2026-11-04)', 'Hudbay Minerals': '待发布(2026-11-12)',
    'Liontown': '待发布(2026-10-30)', 'Sigma Lithium': '待发布(预计2026-11中旬)',
    'SQM': '待发布(预计2026-11中旬)', 'Lithium Argentina': '待发布(预计2026-11上旬)',
    'Elevra Lithium': '待发布(预计2026-11上旬)', 'AMG Critical Materials': '待发布(预计2026-11上旬)',
    'Minsur': '待发布(预计2026-11中旬)', '明苏尔': '待发布(预计2026-11中旬)', 'PT Timah': '待发布(预计2026-11中旬)',
    '紫金矿业': '待发布(预计2026-10月末)', '洛阳钼业': '待发布(预计2026-10月末)',
    '云南锡业': '待发布(预计2026-10月末)', '华锡有色': '待发布(预计2026-10月末)', '兴业银锡': '待发布(预计2026-10月末)',
}
NO_Q3 = {'江西铜业', '铜陵有色', '云南铜业', '华友钴业', '中伟股份', '格林美', '力勤资源',
         '赣锋锂业', '中矿资源', '盛新锂能', '雅化集团', '盐湖股份', '藏格矿业', '西藏矿业',
         '融捷股份', '永兴材料', '江特电机', '天齐锂业', '天华新能', 'MSC'}
SOURCES = {
    '铜': 'Freeport 8-K EX-99.1 10/27；Rio/BHP/Anglo/Antofagasta/MMG/South32官方Q3；Nornickel 9M官方PDF 10/26；First Quantum 10/28；其余公司披露日历',
    '镍': 'Nornickel 9M官方PDF 10/26；Anglo Q3 Production Report 10/20；ANTAM 9M26 10/28；Nickel Industries Q3 Activities Report 10/22；其余公司披露日历',
    '锂': 'PLS Q1 FY27 10/28；MinRes Q1 FY27 10/23；IGO Q1 FY27 10/29；Rio Q3 Operations Review 10/14；Core Finniss restart 9/8；其余公司披露日历',
    '锡': 'Alphamin Q3 update 10/6；Metals X September quarterly 10/28；Andrada operational update 9/28；其余公司披露日历',
}
LOGTEXT = {
    '铜': '开26Q3列并首批入库：FreePort 377.395kt、Antofagasta 151.9、Rio 205、Nornickel本部89.704+外贝加尔19、South32 18.7、BHP 400.7、FQM 108.959、SCCO 232.396、Anglo 174.9、Sandfire 38.9、Ivanhoe权益30.2548、MMG 13.1215万吨；其余按披露日待发布',
    '镍': '开26Q3列并首批入库：Nornickel 61,103、Anglo 10,400、Antam 3,850、NIC RKEF 27,822吨；其余按披露日待发布/无季度披露',
    '锂': '开26Q3列并首批入库：PLS 33,450、MinRes既有Wodgina+Mt Marion 21,625、IGO 54,375、Rio 13,000吨LCE；Core首批复产但官方吨位未披露；其余按披露日待发布/无季度披露',
    '锡': '开26Q3列并首批入库：Alphamin 5,030、Metals X 2,663、Andrada 290吨；其余按披露日待发布/无季度吨数',
}


def ensure_q3(ws):
    headers = [c.value for c in ws[1]]
    if '26Q3' in headers:
        q3_col = headers.index('26Q3') + 1
        q3yoy_col = headers.index('26Q3同比') + 1 if '26Q3同比' in headers else q3_col + 1
        ws.cell(1, q3yoy_col).value = '26Q3同比'
        return q3_col, q3yoy_col
    reason_col = headers.index('变化原因') + 1
    q2_col = headers.index('26Q2') + 1
    q2yoy_col = headers.index('26Q2同比') + 1
    q3_col, q3yoy_col = reason_col, reason_col + 1
    ws.insert_cols(reason_col, 2)
    for row in range(1, ws.max_row + 1):
        src, src_yoy = ws.cell(row, q2_col), ws.cell(row, q2yoy_col)
        dst, dst_yoy = ws.cell(row, q3_col), ws.cell(row, q3yoy_col)
        for target, source in ((dst, src), (dst_yoy, src_yoy)):
            if source.has_style:
                target._style = copy.copy(source._style)
            if source.number_format:
                target.number_format = source.number_format
        if row == 1:
            dst.value, dst_yoy.value = '26Q3', '26Q3同比'
            continue
        if src.value is not None and isinstance(src.value, str) and src.value.startswith('='):
            dst.value = Translator(src.value, origin=src.coordinate).translate_formula(dst.coordinate)
        if ws.cell(row, 1).value == '总计':
            dst.value = Translator(src.value, origin=src.coordinate).translate_formula(dst.coordinate)
        prior_q3 = ws.cell(row, headers.index('25Q3') + 1).coordinate
        dst_yoy.value = f'=IF(AND(ISNUMBER({dst.coordinate}),ISNUMBER({prior_q3}),{prior_q3}<>0),({dst.coordinate}-{prior_q3})/{prior_q3},"-")'
    ws.column_dimensions[openpyxl.utils.get_column_letter(q3_col)].width = ws.column_dimensions[openpyxl.utils.get_column_letter(q2_col)].width
    ws.column_dimensions[openpyxl.utils.get_column_letter(q3yoy_col)].width = ws.column_dimensions[openpyxl.utils.get_column_letter(q2yoy_col)].width
    return q3_col, q3yoy_col


def update_workbook(comm, path):
    backup = path.with_name(path.stem + '_backup_20261029.xlsx')
    if not backup.exists():
        shutil.copy2(path, backup)
    wb = openpyxl.load_workbook(path)
    confirmed = {(s, c, p): (v, r) for s, c, p, v, r in DATA[comm]}
    touched = []
    for ws in wb:
        if '日志' in ws.title:
            continue
        if '26Q2' not in [c.value for c in ws[1]]:
            continue
        q3_col, _ = ensure_q3(ws)
        reason_col = [c.value for c in ws[1]].index('变化原因') + 1
        q2_col = [c.value for c in ws[1]].index('26Q2') + 1
        nature_col = [c.value for c in ws[1]].index('数据性质') + 1 if '数据性质' in [c.value for c in ws[1]] else None
        for row in range(2, ws.max_row + 1):
            company = ws.cell(row, 1).value
            if not company or company in ('总计',) or str(company).startswith('注：'):
                continue
            project = str(ws.cell(row, 3).value or '')
            match = next(((v, r) for (s, c, p), (v, r) in confirmed.items()
                          if s == ws.title and c == company and (p is None or p in project)), None)
            if match:
                value, reason = match
            elif company in PENDING:
                value, reason = PENDING[company], '截至2026-10-29核查，按公司披露日历等待官方产量'
            elif company in NO_Q3:
                value, reason = '-', 'A股/港股季报无对应产量吨数或口径不符，保持空缺'
            elif ws.cell(row, q2_col).value in (None, '', '-'):
                value, reason = '-', '无季度披露或既有口径不适用'
            else:
                value, reason = '待发布(预计2026-11上旬)', '截至2026-10-29核查，等待官方产量披露'
            ws.cell(row, q3_col).value = value
            ws.cell(row, reason_col).value = reason
            if nature_col:
                ws.cell(row, nature_col).value = '披露' if isinstance(value, (int, float)) else '状态'
            touched.append((ws.title, company, value))
    for ws in wb:
        if '日志' in ws.title:
            last = [cell.value for cell in ws[ws.max_row]]
            if last[:3] != [TODAY, LOGTEXT[comm], SOURCES[comm]]:
                ws.append([TODAY, LOGTEXT[comm], SOURCES[comm]])
            break
    wb.save(path)
    return touched, backup


def update_disruptions():
    path = Path(r'D:/Kimi/caibao-hub/data/disruptions.json')
    data = json.loads(path.read_text(encoding='utf-8'))
    for item in data['items']:
        if item['company'] == 'FreePort Grasberg':
            item['impact'] = '2025/9事故后爬产，26Q3磨矿处理量14.1万吨/日=正常产能67%；集团Q3铜832M lbs'
            item['recovery'] = 'PTFI冶炼厂2026年8月底重启；GBC预计2027年中80%、2027年底接近满产；2027为恢复大年'
            item['source'] = 'FreePort 3Q26 8-K EX-99.1（2026-10-27）'
            item['url'] = 'https://www.sec.gov/Archives/edgar/data/831259/000083125926000045/a3q2026exhibit991.htm'
        if item['company'] == 'Weda Bay（Eramet/青山）':
            item['recovery'] = '印尼政府授权后，PT WBN于2026-09-10宣布矿山恢复运营；后续产量与配额执行仍待Q3/Q4披露验证'
            item['source'] = 'Eramet：PT Weda Bay Nickel restarts its mining operations（2026-09-10）'
            item['url'] = 'https://www.eramet.com/en/news/pt-weda-bay-nickel-restarts-its-mining-operations/'
        if item['company'] == 'Cobre Panamá 堆存矿获批（First Quantum）':
            item['impact'] = '重启调试取得实质进展，Q3三条产线已陆续恢复运行；集团Q3铜108,959吨'
            item['recovery'] = '三条产线恢复运行；政府间委员会建议获公司积极评价，持久解决方案仍待与政府推进'
            item['source'] = 'First Quantum Q3 2026 Results（2026-10-28）'
            item['url'] = 'https://www.first-quantum.com/news/first-quantum-minerals-reports-third-quarter-2026-results/'
    additions = [
        {'commodity':'镍','date':'2026-10-22','ongoing':True,'company':'Nickel Industries RKEF / ENC HPAL','country':'印尼','type':'天气·干旱缺水','dir':'down','capacity':'RKEF约12.5万吨镍/年；ENC HPAL 7.2万吨镍/年','impact':'Hengjaya地区8月至10月下旬降雨仅7mm（十年均值474mm）；Q3 RKEF镍27,822吨，10月仅约75%铭牌负荷，ENC爬产亦受限','recovery':'等待雨季回归，管理层预计12月季缓解；恢复以实际降雨、RKEF负荷和ENC产量验证','source':'Nickel Industries Q3 2026 Quarterly Activities Report','url':'https://nickelindustries.com/'},
        {'commodity':'铜','date':'2026-10-20','company':'Anglo American（智利铜矿）','country':'智利','type':'天气·创纪录降雪','dir':'down','capacity':'2026铜指引700-760kt','impact':'Q3铜174.9kt（同比-5%）：智利暴雪及非计划检修，Collahuasi处理低品位库存矿','recovery':'Q4计划品位改善，全年指引维持；需以Q4实际产量验证恢复','source':'Anglo American Q3 2026 Production Report','url':'https://www.angloamerican.com/media/press-releases/2026/20-10-2026'},
        {'commodity':'铜','date':'2026-10-20','company':'BHP（智利/Carrapateena）','country':'智利/澳大利亚','type':'天气+设备故障','dir':'down','capacity':'FY27铜指引1650-1800kt','impact':'FY27Q1铜400.7kt（同比-19%）：计划品位下降与检修，叠加暴雪、安全停产和Carrapateena输送带故障','recovery':'天气与设备为阶段性扰动，品位下降为结构性；后续看Q2 FY27修复','source':'BHP Operational Review September 2026','url':'https://www.bhp.com/news/media-centre/releases/2026/10/bhp-operational-review-for-the-quarter-ended-30-september-2026'},
        {'commodity':'锌','date':'2026-10-19','company':'South32 Cannington','country':'澳大利亚','type':'设备·回填问题','dir':'down','capacity':'FY27锌当量指引20.47万吨','impact':'临时paste backfill问题导致九月季度产量同比-27%；Q3应付锌0.88万吨、铅1.6万吨','recovery':'公司维持FY27指引，后续看回填修复和Q2产量恢复','source':'South32 September 2026 Quarterly Report','url':'https://www.south32.net/news-media/latest-news/quarterly-report-september-2026'},
        {'commodity':'锌','date':'2026-10-14','company':'MMG Rosebery（X-South）','country':'澳大利亚','type':'事故·地震','dir':'down','capacity':'原2026锌指引4.5-5.5万吨','impact':'X-South地震影响，Rosebery 2026锌指引下调至3.2-3.7万吨；Q3锌0.813万吨、铅0.2969万吨','recovery':'按新指引和矿山恢复进度跟踪，暂不假设年内恢复至原指引','source':'MMG Q3 2026 Production Report','url':'https://www.mmg.com/content/uploads/2026/10/e_2026-10-14_3QTR-Production-Report.pdf'},
        {'commodity':'锌','date':'2026-10-08','company':'Ivanhoe Kipushi','country':'刚果(金)','type':'爬产·季度纪录','dir':'up','capacity':'高品位锌矿','impact':'Q3锌精矿77,147吨（环比+10%）创季度纪录','recovery':'持续爬产，后续看稳态运行与回收率','source':'Ivanhoe Mines Q3 2026 Production Update','url':'https://www.ivanhoemines.com/'},
        {'commodity':'锡','date':'2026-10-06','company':'Alphamin Bisie','country':'刚果(金)','type':'天气·雨季物流风险','dir':'down','capacity':'约2万吨锡/年','impact':'Q3含锡5,030吨持稳；El Niño或令Q4雨季运输更具扰动，属于前瞻风险而非已发生减产','recovery':'已预增卡车运力与燃料储备；Q4发运和道路状况待验证','source':'Alphamin Q3 2026 Operational Update','url':'https://www.alphaminresources.com/2026/10/06/alphamin-announces-cad0-13-per-share-interim-fy2026-dividend-q3-guidance-exploration-update/'},
        {'commodity':'锂','date':'2026-09-08','company':'Core Lithium Finniss','country':'澳大利亚','type':'复产·首批精矿','dir':'up','capacity':'重启目标约214ktpa精矿','impact':'重启后首批锂精矿产出，Grants露天矿供料；官方尚未披露Q3吨位','recovery':'Q4目标首批发运，BP33地下开发继续，后续看稳定产量','source':'Core Lithium Finniss restart update','url':'https://www.corelithium.com.au/'},
        {'commodity':'锂','date':'2026-10-23','company':'MinRes Bald Hill','country':'澳大利亚','type':'复产爬坡','dir':'up','capacity':'设计约140ktpa SC6（Q2 FY27目标）','impact':'FY27Q1精矿27k dmt、销量21k dmt SC6，回收率63%，重启增量开始兑现','recovery':'DMS厂向设计产能爬坡，FY27成本与产量仍处爬坡期','source':'MinRes Q1 FY27 Quarterly Activity Report','url':'https://www.mineralresources.com.au/news/company-updates/q1-fy27-quarterly-activity-report/'},
        {'commodity':'锂','date':'2026-10-29','company':'IGO Greenbushes（CGP3）','country':'澳大利亚','type':'扩产·季度纪录','dir':'up','capacity':'CGP3扩产项目','impact':'FY27Q1精矿435kt（环比+12%）创纪录，CGP3达到铭牌产能；折LCE 54,375吨','recovery':'扩产已兑现，后续关注矿石品位、回收率与稳定运行','source':'IGO Q1 FY27 Quarterly Activities Report','url':'https://www.igo.com.au/'},
    ]
    existing = {(x.get('company'), x.get('date')) for x in data['items']}
    for item in additions:
        if (item['company'], item['date']) not in existing:
            data['items'].append(item)
    data['updated'] = TODAY
    backup = path.with_name('disruptions_backup_20261029.json')
    if not backup.exists():
        shutil.copy2(path, backup)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')
    return len(additions)


if __name__ == '__main__':
    for comm, path in FILES.items():
        touched, backup = update_workbook(comm, path)
        numeric = sum(isinstance(v, (int, float)) for _s, _c, v in touched)
        print(comm, '26Q3数值行', numeric, '备份', backup)
    print('供应扰动新增候选', update_disruptions())
